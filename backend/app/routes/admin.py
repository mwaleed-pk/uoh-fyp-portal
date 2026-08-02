"""Secure administration APIs for reporting and account management."""
from datetime import datetime, timezone
from io import BytesIO, StringIO
import csv
from typing import Literal, Optional

from fastapi import APIRouter, Depends, File, HTTPException, Query, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, EmailStr, Field, TypeAdapter, ValidationError

from app.db.supabase_client import get_supabase_client
from app.services.notification_service import create_notification
from app.utils.audit import write_audit_log
from app.utils.auth_middleware import require_admin

router = APIRouter()

STUDENT_FIELDS = (
    "id,name,email,phone,student_number,registration_number,profile_photo_url,status,"
    "department_id,program_id,semester_id,session_id,section_id,created_at,updated_at,"
    "departments(name,code),programs(name,code),semesters(name),sessions(academic_year),sections(name),"
    "projects(id,title,idea,abstract,status,lifecycle_stage,supervisor_id,supervisors(id,name,email))"
)
SUPERVISOR_FIELDS = (
    "id,name,email,phone,profile_photo_url,department_id,designation,max_students,areas_of_interest,"
    "status,approval_status,rejection_note,email_verified_at,created_at,updated_at,"
    "departments(name,code),supervisor_program_eligibility(program_id,programs(name,code))"
)


class AccountAction(BaseModel):
    action: Literal["activate", "deactivate", "approve", "reject", "revoke_sessions"]
    note: Optional[str] = Field(default=None, max_length=1000)


class AssignmentBody(BaseModel):
    supervisor_id: Optional[str] = None
    reason: Optional[str] = Field(default=None, max_length=1000)
    enforce_capacity: bool = True

class StudentUpdate(BaseModel):
    name: Optional[str] = None
    email: Optional[EmailStr] = None
    phone: Optional[str] = None
    student_number: Optional[str] = None
    registration_number: Optional[str] = None
    profile_photo_url: Optional[str] = None
    department_id: Optional[str] = None
    program_id: Optional[str] = None
    semester_id: Optional[str] = None
    session_id: Optional[str] = None
    section_id: Optional[str] = None
    status: Optional[Literal["active", "inactive", "suspended"]] = None
    project_title: Optional[str] = None
    project_idea: Optional[str] = None

class StudentCreate(BaseModel):
    name: str = Field(min_length=2, max_length=150)
    email: EmailStr
    phone: Optional[str] = Field(default=None, max_length=40)
    student_number: str = Field(min_length=1, max_length=64)
    registration_number: str = Field(min_length=1, max_length=100)
    department_id: str
    program_id: str
    semester_id: Optional[str] = None
    session_id: str
    section_id: str
    profile_photo_url: Optional[str] = Field(default=None, max_length=2000)
    status: Literal["active", "inactive"] = "active"
    supervisor_id: Optional[str] = None

class SupervisorUpdate(BaseModel):
    name: Optional[str] = None
    email: Optional[EmailStr] = None
    phone: Optional[str] = None
    profile_photo_url: Optional[str] = None
    department_id: Optional[str] = None
    designation: Optional[str] = None
    areas_of_interest: Optional[str] = None
    max_students: Optional[int] = Field(default=None, ge=1, le=100)
    status: Optional[Literal["active", "inactive", "suspended"]] = None
    program_ids: Optional[list[str]] = None

class SupervisorCreate(BaseModel):
    name: str = Field(min_length=2, max_length=150)
    email: EmailStr
    phone: Optional[str] = Field(default=None, max_length=40)
    department_id: str
    designation: Optional[str] = Field(default=None, max_length=150)
    areas_of_interest: Optional[str] = Field(default=None, max_length=2000)
    profile_photo_url: Optional[str] = Field(default=None, max_length=2000)
    status: Literal["active", "inactive"] = "active"
    program_ids: list[str] = Field(default_factory=list)

class WorkflowDecision(BaseModel):
    decision: Literal["approved", "rejected"]
    note: Optional[str] = Field(default=None, max_length=1000)


def _count(table: str, **equals) -> int:
    q = get_supabase_client().table(table).select("id", count="exact")
    for key, value in equals.items():
        q = q.eq(key, value)
    result = q.execute()
    return result.count or 0


def _require_active_reference(client, table: str, value: Optional[str], label: str):
    if not value:
        return
    if not client.table(table).select("id").eq("id", value).eq("status", "active").execute().data:
        raise HTTPException(422, f"Selected {label} is unavailable.")


def _validate_student_relationships(client, data: dict):
    for field, table, label in (
        ("department_id", "departments", "department"), ("program_id", "programs", "program"),
        ("semester_id", "semesters", "semester"), ("session_id", "sessions", "academic session"),
        ("section_id", "sections", "section"),
    ):
        _require_active_reference(client, table, data.get(field), label)


def _normalized(value) -> str:
    return str(value or "").strip()


def _assign_supervisor(client, project_id: str, supervisor_id: Optional[str], admin_id: str,
                       reason: Optional[str]):
    try:
        result = client.rpc("assign_project_supervisor", {"p_project_id": project_id,
                            "p_supervisor_id": supervisor_id, "p_changed_by": admin_id,
                            "p_reason": reason}).execute()
        return (result.data or [{}])[0]
    except Exception as exc:
        message = str(exc)
        if "CAPACITY_FULL" in message: raise HTTPException(409, "Supervisor capacity is full.") from exc
        if "SUPERVISOR_UNAVAILABLE" in message: raise HTTPException(422, "Supervisor must be active and approved.") from exc
        if "PROJECT_NOT_FOUND" in message: raise HTTPException(404, "Student project not found.") from exc
        raise


SUPERVISOR_IMPORT_COLUMNS = ["name", "email", "phone", "department", "designation", "areas_of_interest", "status"]


def _supervisor_template_rows():
    return [{"name": "Dr Example", "email": "faculty@uoh.edu.pk", "phone": "03000000000",
             "department": "CS", "designation": "Assistant Professor",
             "areas_of_interest": "Artificial Intelligence", "status": "active"}]


def _safe_spreadsheet_cell(value):
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def _tabular_response(rows: list[dict], columns: list[str], fmt: str, filename: str):
    rows = [{column: _safe_spreadsheet_cell(row.get(column)) for column in columns} for row in rows]
    if fmt == "csv":
        out = StringIO(); writer = csv.DictWriter(out, fieldnames=columns); writer.writeheader(); writer.writerows(rows)
        return Response(("\ufeff" + out.getvalue()).encode("utf-8"), media_type="text/csv; charset=utf-8",
                        headers={"Content-Disposition": f'attachment; filename="{filename}.csv"', "X-Content-Type-Options": "nosniff"})
    wb = Workbook(); ws = wb.active; ws.title = "Supervisors"; ws.append(columns)
    for row in rows: ws.append([row.get(column) for column in columns])
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(45, max(12, max(len(str(c.value or "")) for c in col) + 2))
    out = BytesIO(); wb.save(out)
    return Response(out.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{filename}.xlsx"', "X-Content-Type-Options": "nosniff"})


@router.get("/dashboard")
async def dashboard(_: dict = Depends(require_admin)):
    client = get_supabase_client()
    import asyncio
    # Execute queries sequentially because the synchronous Supabase client
    # uses a single httpx connection pool which can deadlock under heavy threading.
    projects_res = client.table("projects").select("id,student_id,supervisor_id,lifecycle_stage,status").execute()
    docs_res = client.table("project_documents").select("id,status,created_at").execute()
    recent_students_res = client.table("students").select("id,name,created_at").order("created_at", desc=True).limit(5).execute()
    recent_supervisors_res = client.table("supervisors").select("id,name,created_at").order("created_at", desc=True).limit(5).execute()
    recent_activity_res = client.table("audit_logs").select("id,actor_id,actor_role,action_type,resource_type,resource_id,metadata,created_at").order("created_at", desc=True).limit(10).execute()
    assigned_res = client.table("projects").select("student_id").not_.is_("supervisor_id", "null").execute()
    
    projects = projects_res.data or []
    docs = docs_res.data or []
    recent_students = recent_students_res.data or []
    recent_supervisors = recent_supervisors_res.data or []
    recent_activity = recent_activity_res.data or []
    assigned = {p["student_id"] for p in (assigned_res.data or [])}
    
    c_students = _count("students")
    c_active_students = _count("students", status="active")
    c_supervisors = _count("supervisors")
    c_active_supervisors = _count("supervisors", status="active")
    c_pending_supervisors = _count("supervisors", approval_status="pending")
    
    counts = {
        "students": c_students, "active_students": c_active_students,
        "supervisors": c_supervisors, "active_supervisors": c_active_supervisors,
        "pending_supervisors": c_pending_supervisors,
        "projects": len(projects), "students_without_supervisors": max(0, c_students - len(assigned)),
        "pending_reviews": sum(1 for d in docs if d.get("status") in ("submitted", "pending")),
        "approved_submissions": sum(1 for d in docs if d.get("status") == "approved"),
        "rejected_submissions": sum(1 for d in docs if d.get("status") in ("rejected", "revision_requested")),
    }
    return {"counts": counts, "total_students": counts["students"], "active_students": counts["active_students"],
            "total_supervisors": counts["supervisors"], "pending_supervisors": counts["pending_supervisors"],
            "total_projects": counts["projects"], "unassigned_students": counts["students_without_supervisors"],
            "pending_reviews": counts["pending_reviews"], "approved_submissions": counts["approved_submissions"],
            "recent_registrations": recent_students + recent_supervisors, "recent_activity": recent_activity}


def _paginate(q, page: int, page_size: int):
    return q.range((page - 1) * page_size, page * page_size - 1)


@router.get("/students")
async def students(search: Optional[str] = None, department_id: Optional[str] = None,
                   program_id: Optional[str] = None, semester_id: Optional[str] = None,
                   session_id: Optional[str] = None, section_id: Optional[str] = None,
                   supervisor_id: Optional[str] = None, status: Optional[str] = None,
                   project_status: Optional[str] = None, sort: str = "created_at", descending: bool = True,
                   page: int = Query(1, ge=1), page_size: int = Query(25, ge=1, le=200),
                   _: dict = Depends(require_admin)):
    client = get_supabase_client()
    q = client.table("students").select(STUDENT_FIELDS, count="exact")
    if supervisor_id or project_status:
        project_q = client.table("projects").select("student_id")
        if supervisor_id: project_q = project_q.eq("supervisor_id", supervisor_id)
        if project_status: project_q = project_q.eq("status", project_status)
        matching_ids = [p["student_id"] for p in (project_q.execute().data or [])]
        if not matching_ids:
            return {"students": [], "count": 0, "page": page, "page_size": page_size}
        q = q.in_("id", matching_ids)
    for key, value in (("department_id", department_id), ("program_id", program_id), ("semester_id", semester_id),
                       ("session_id", session_id), ("section_id", section_id), ("status", status)):
        if value: q = q.eq(key, value)
    if search:
        term = search.replace(",", " ").strip()
        q = q.or_(f"name.ilike.%{term}%,email.ilike.%{term}%,student_number.ilike.%{term}%,registration_number.ilike.%{term}%")
    safe_sort = sort if sort in {"name", "email", "student_number", "created_at", "updated_at", "status"} else "created_at"
    result = _paginate(q.order(safe_sort, desc=descending), page, page_size).execute()
    total = result.count or 0
    return {"students": result.data or [], "count": total, "total": total,
            "pages": max(1, (total + page_size - 1) // page_size), "page": page, "page_size": page_size}


@router.get("/students/{student_id}")
async def student_detail(student_id: str, _: dict = Depends(require_admin)):
    client = get_supabase_client()
    row = client.table("students").select(STUDENT_FIELDS).eq("id", student_id).single().execute().data
    if not row: raise HTTPException(404, "Student not found.")
    project_ids = [p["id"] for p in row.get("projects") or []]
    row["documents"] = client.table("project_documents").select("id,project_id,filename,document_category,status,version_number,created_at,updated_at").in_("project_id", project_ids).order("created_at", desc=True).execute().data if project_ids else []
    return {"student": row}

@router.post("/students", status_code=201)
async def create_student(body: StudentCreate, admin: dict = Depends(require_admin)):
    client = get_supabase_client(); data = body.model_dump()
    data["email"] = str(body.email).strip().lower()
    data["name"] = body.name.strip(); data["student_number"] = body.student_number.strip()
    data["registration_number"] = body.registration_number.strip()
    supervisor_id = data.pop("supervisor_id", None)
    _validate_student_relationships(client, data)
    duplicate = client.table("students").select("id,email,student_number,registration_number").or_(
        f"email.eq.{data['email']},student_number.eq.{data['student_number']},registration_number.eq.{data['registration_number']}"
    ).execute().data or []
    if duplicate: raise HTTPException(409, "A student with this email, roll number, or registration number already exists.")
    if supervisor_id:
        sup = client.table("supervisors").select("id,status,approval_status").eq("id", supervisor_id).single().execute().data
        if not sup or sup.get("status") != "active" or sup.get("approval_status") != "approved":
            raise HTTPException(422, "Selected supervisor is unavailable.")
    data["created_by"] = admin["user_id"]
    created = client.table("students").insert(data).execute().data or []
    if not created: raise HTTPException(500, "Student could not be created.")
    student = created[0]
    project = client.table("projects").insert({"student_id": student["id"], "supervisor_id": None}).execute().data or []
    try:
        if supervisor_id and project:
            _assign_supervisor(client, project[0]["id"], supervisor_id, admin["user_id"], "Initial admin assignment")
    except Exception:
        if project: client.table("projects").delete().eq("id", project[0]["id"]).execute()
        client.table("students").delete().eq("id", student["id"]).execute()
        raise
    await write_audit_log("account_created", admin["user_id"], "admin", "student", student["id"])
    return {"message": "Student created.", "student": student, "project": project[0] if project else None}

@router.patch("/students/{student_id}")
async def update_student(student_id: str, body: StudentUpdate, admin: dict = Depends(require_admin)):
    client=get_supabase_client(); previous=client.table("students").select("id,name,phone,department_id,program_id,semester_id,session_id,section_id,status").eq("id",student_id).single().execute().data
    if not previous: raise HTTPException(404,"Student not found.")
    data=body.model_dump(exclude_none=True); project_data={}
    if "email" in data: data["email"] = str(data["email"]).strip().lower()
    _validate_student_relationships(client, data)
    if "project_title" in data: project_data["title"]=data.pop("project_title")
    if "project_idea" in data: project_data["idea"]=data.pop("project_idea")
    if data: client.table("students").update(data).eq("id",student_id).execute()
    if project_data: client.table("projects").update(project_data).eq("student_id",student_id).execute()
    if data.get("status") in {"inactive","suspended"}: client.table("user_sessions").delete().eq("user_id",student_id).eq("user_role","student").execute()
    await write_audit_log("account_updated",admin["user_id"],"admin","student",student_id,metadata={"previous":previous,"new":data|project_data})
    return {"message":"Student updated."}


@router.get("/supervisors")
async def supervisors(search: Optional[str] = None, department_id: Optional[str] = None,
                      program_id: Optional[str] = None, status: Optional[str] = None,
                      approval_status: Optional[str] = None, page: int = Query(1, ge=1),
                      page_size: int = Query(25, ge=1, le=200), _: dict = Depends(require_admin)):
    client = get_supabase_client(); q = client.table("supervisors").select(SUPERVISOR_FIELDS, count="exact")
    for key, value in (("department_id", department_id), ("status", status), ("approval_status", approval_status)):
        if value: q = q.eq(key, value)
    if search:
        term = search.replace(",", " ").strip(); q = q.or_(f"name.ilike.%{term}%,email.ilike.%{term}%,designation.ilike.%{term}%")
    result = _paginate(q.order("created_at", desc=True), page, page_size).execute(); data = result.data or []
    if program_id: data = [s for s in data if any(e.get("program_id") == program_id for e in s.get("supervisor_program_eligibility") or [])]
    for s in data:
        projects = client.table("projects").select("id,lifecycle_stage").eq("supervisor_id", s["id"]).execute().data or []
        s["assigned_student_count"] = len(projects); s["active_project_count"] = sum(p.get("lifecycle_stage") != "closed" for p in projects)
    total = result.count or 0
    return {"supervisors": data, "count": total, "total": total,
            "pages": max(1, (total + page_size - 1) // page_size), "page": page, "page_size": page_size}


@router.get("/supervisors/import-template")
async def supervisor_import_template(format: Literal["csv", "xlsx"] = "xlsx",
                                     _: dict = Depends(require_admin)):
    return _tabular_response(_supervisor_template_rows(), SUPERVISOR_IMPORT_COLUMNS, format,
                             "supervisor-import-template")


@router.get("/supervisors/{supervisor_id}")
async def supervisor_detail(supervisor_id: str, _: dict = Depends(require_admin)):
    client = get_supabase_client(); row = client.table("supervisors").select(SUPERVISOR_FIELDS).eq("id", supervisor_id).single().execute().data
    if not row: raise HTTPException(404, "Supervisor not found.")
    row["projects"] = client.table("projects").select("id,title,status,lifecycle_stage,student_id,students(id,name,email,student_number,programs(name),semesters(name),sections(name))").eq("supervisor_id", supervisor_id).execute().data or []
    return {"supervisor": row}

@router.post("/supervisors", status_code=201)
async def create_supervisor(body: SupervisorCreate, admin: dict = Depends(require_admin)):
    client = get_supabase_client(); data = body.model_dump()
    program_ids = list(dict.fromkeys(data.pop("program_ids", [])))
    data["email"] = str(body.email).strip().lower(); data["name"] = body.name.strip()
    _require_active_reference(client, "departments", data["department_id"], "department")
    for program_id in program_ids: _require_active_reference(client, "programs", program_id, "program")
    if client.table("supervisors").select("id").eq("email", data["email"]).execute().data:
        raise HTTPException(409, "A supervisor with this email already exists.")
    data.update({"created_by": admin["user_id"], "approval_status": "approved", "max_students": 30})
    created = client.table("supervisors").insert(data).execute().data or []
    if not created: raise HTTPException(500, "Supervisor could not be created.")
    supervisor = created[0]
    if program_ids:
        client.table("supervisor_program_eligibility").insert([{"supervisor_id": supervisor["id"], "program_id": p} for p in program_ids]).execute()
    await write_audit_log("account_created", admin["user_id"], "admin", "supervisor", supervisor["id"])
    return {"message": "Supervisor created.", "supervisor": supervisor}

@router.patch("/supervisors/{supervisor_id}")
async def update_supervisor(supervisor_id: str, body: SupervisorUpdate, admin: dict = Depends(require_admin)):
    client=get_supabase_client(); previous=client.table("supervisors").select("id,name,phone,department_id,designation,areas_of_interest,max_students,status").eq("id",supervisor_id).single().execute().data
    if not previous: raise HTTPException(404,"Supervisor not found.")
    data=body.model_dump(exclude_none=True); program_ids=data.pop("program_ids",None)
    if "email" in data: data["email"] = str(data["email"]).strip().lower()
    _require_active_reference(client,"departments",data.get("department_id"),"department")
    if program_ids is not None:
        program_ids=list(dict.fromkeys(program_ids))
        for program_id in program_ids: _require_active_reference(client,"programs",program_id,"program")
    if data: client.table("supervisors").update(data).eq("id",supervisor_id).execute()
    if program_ids is not None:
        client.table("supervisor_program_eligibility").delete().eq("supervisor_id",supervisor_id).execute()
        if program_ids: client.table("supervisor_program_eligibility").insert([{"supervisor_id":supervisor_id,"program_id":p} for p in program_ids]).execute()
    if data.get("status") in {"inactive","suspended"}: client.table("user_sessions").delete().eq("user_id",supervisor_id).eq("user_role","supervisor").execute()
    await write_audit_log("account_updated",admin["user_id"],"admin","supervisor",supervisor_id,metadata={"previous":previous,"new":data})
    return {"message":"Supervisor updated."}


@router.delete("/{entity}/{user_id}")
async def safely_deactivate(entity: Literal["students", "supervisors"], user_id: str,
                            admin: dict = Depends(require_admin)):
    role = "student" if entity == "students" else "supervisor"; client = get_supabase_client()
    row = client.table(entity).select("id,status").eq("id", user_id).single().execute().data
    if not row: raise HTTPException(404, f"{role.title()} not found.")
    client.table(entity).update({"status": "inactive"}).eq("id", user_id).execute()
    client.table("user_sessions").delete().eq("user_id", user_id).eq("user_role", role).execute()
    await write_audit_log("account_deactivated", admin["user_id"], "admin", role, user_id,
                          metadata={"previous_status": row.get("status"), "safe_delete": True})
    return {"message": f"{role.title()} deactivated safely.", "status": "inactive"}


@router.post("/{role}/{user_id}/account-action")
async def account_action(role: Literal["student", "supervisor"], user_id: str, body: AccountAction,
                         admin: dict = Depends(require_admin)):
    client = get_supabase_client(); table = "students" if role == "student" else "supervisors"
    previous = client.table(table).select("id,status" + (",approval_status,rejection_note" if role == "supervisor" else "")).eq("id", user_id).single().execute().data
    if not previous: raise HTTPException(404, f"{role.title()} not found.")
    update = {}
    if body.action == "activate": update["status"] = "active"
    elif body.action == "deactivate": update["status"] = "inactive"
    elif body.action in ("approve", "reject"):
        if role != "supervisor": raise HTTPException(422, "Approval applies only to supervisors.")
        if body.action == "reject" and not (body.note or "").strip(): raise HTTPException(422, "A rejection note is required.")
        update.update({"approval_status": "approved" if body.action == "approve" else "rejected", "rejection_note": body.note if body.action == "reject" else None})
        if body.action == "approve": update["status"] = "active"
    if update: client.table(table).update(update).eq("id", user_id).execute()
    if body.action in ("deactivate", "revoke_sessions", "reject"):
        client.table("user_sessions").delete().eq("user_id", user_id).eq("user_role", role).execute()
    await write_audit_log("account_deactivated" if body.action == "deactivate" else "account_updated", admin["user_id"], "admin", role, user_id, metadata={"action": body.action, "previous": previous, "new": update, "note": body.note})
    return {"message": f"Account action '{body.action}' completed."}


@router.put("/students/{student_id}/assignment")
async def set_assignment(student_id: str, body: AssignmentBody, admin: dict = Depends(require_admin)):
    client = get_supabase_client(); project = client.table("projects").select("id,supervisor_id,status,lifecycle_stage").eq("student_id", student_id).single().execute().data
    if not project: raise HTTPException(404, "Student project not found.")
    old_id = project.get("supervisor_id")
    assignment = _assign_supervisor(client, project["id"], body.supervisor_id, admin["user_id"], body.reason)
    if body.supervisor_id:
        await create_notification(student_id, "student", "supervisor_assigned", "Supervisor assignment updated", "Your supervisor assignment has been updated.", project["id"])
        await create_notification(body.supervisor_id, "supervisor", "supervisor_assigned", "New student assigned", "A student has been assigned to you.", project["id"])
    await write_audit_log("supervisor_assigned", admin["user_id"], "admin", "project", project["id"], metadata={"previous_supervisor_id": old_id, "new_supervisor_id": body.supervisor_id, "reason": body.reason})
    return {"message": "Assignment updated.", "assignment": assignment}


@router.get("/audit-logs")
async def audit_logs(actor_id: Optional[str] = None, action_type: Optional[str] = None,
                     resource_type: Optional[str] = None, page: int = Query(1, ge=1),
                     page_size: int = Query(50, ge=1, le=200), _: dict = Depends(require_admin)):
    q = get_supabase_client().table("audit_logs").select("id,actor_id,actor_role,action_type,resource_type,resource_id,metadata,ip_address,created_at", count="exact")
    for key, value in (("actor_id", actor_id), ("action_type", action_type), ("resource_type", resource_type)):
        if value: q = q.eq(key, value)
    result = _paginate(q.order("created_at", desc=True), page, page_size).execute()
    return {"logs": result.data or [], "count": result.count or 0, "page": page, "page_size": page_size}


@router.post("/academic-change-requests/{request_id}/decision")
async def decide_academic_change(request_id: str, body: WorkflowDecision, admin: dict = Depends(require_admin)):
    client=get_supabase_client(); req=client.table("student_academic_change_requests").select("*").eq("id",request_id).single().execute().data
    if not req: raise HTTPException(404,"Request not found.")
    if req.get("status") != "pending": raise HTTPException(409,"Request has already been decided.")
    if body.decision == "rejected" and not (body.note or "").strip(): raise HTTPException(422,"A rejection note is required.")
    if body.decision == "approved":
        allowed={"department_id","program_id","semester_id","session_id","section_id"}; changes={k:v for k,v in (req.get("requested_changes") or {}).items() if k in allowed}
        if not changes: raise HTTPException(422,"Request contains no valid academic changes.")
        tables={"department_id":"departments","program_id":"programs","semester_id":"semesters","session_id":"sessions","section_id":"sections"}
        for field,value in changes.items():
            check=client.table(tables[field]).select("id").eq("id",value)
            if field in {"department_id","program_id","semester_id","session_id","section_id"}: check=check.eq("status","active")
            if not check.execute().data: raise HTTPException(422,f"Requested {field.replace('_id','')} is unavailable.")
        client.table("students").update(changes).eq("id",req["student_id"]).execute()
    client.table("student_academic_change_requests").update({"status":body.decision,"decision_note":body.note,"decided_by":admin["user_id"],"decided_at":datetime.now(timezone.utc).isoformat()}).eq("id",request_id).execute()
    await write_audit_log("account_updated",admin["user_id"],"admin","academic_change_request",request_id,metadata={"decision":body.decision,"student_id":req["student_id"]})
    return {"message":f"Academic change request {body.decision}."}


@router.post("/supervisor-reassignment-requests/{request_id}/decision")
async def decide_reassignment(request_id: str, body: WorkflowDecision, admin: dict = Depends(require_admin)):
    client=get_supabase_client(); req=client.table("supervisor_reassignment_requests").select("*").eq("id",request_id).single().execute().data
    if not req: raise HTTPException(404,"Request not found.")
    if req.get("status") != "pending": raise HTTPException(409,"Request has already been decided.")
    if body.decision == "rejected" and not (body.note or "").strip(): raise HTTPException(422,"A rejection note is required.")
    if body.decision == "approved":
        target=req.get("requested_supervisor_id")
        if not target: raise HTTPException(422,"Select a requested supervisor before approval.")
        _assign_supervisor(client, req["project_id"], target, admin["user_id"], body.note or req.get("reason"))
    client.table("supervisor_reassignment_requests").update({"status":body.decision,"decision_note":body.note,"decided_by":admin["user_id"],"decided_at":datetime.now(timezone.utc).isoformat()}).eq("id",request_id).execute()
    await write_audit_log("supervisor_assigned" if body.decision=="approved" else "account_updated",admin["user_id"],"admin","supervisor_reassignment_request",request_id,metadata={"decision":body.decision,"student_id":req["student_id"]})
    return {"message":f"Supervisor reassignment request {body.decision}."}


@router.post("/supervisors/import")
async def import_supervisors(file: UploadFile = File(...), admin: dict = Depends(require_admin)):
    filename = (file.filename or "").lower(); extension = filename.rsplit(".", 1)[-1] if "." in filename else ""
    if extension not in {"csv", "xlsx"}: raise HTTPException(415, "Upload a CSV or XLSX file.")
    payload = await file.read(5 * 1024 * 1024 + 1)
    if not payload: raise HTTPException(422, "The uploaded file is empty.")
    if len(payload) > 5 * 1024 * 1024: raise HTTPException(413, "Import files must be 5 MB or smaller.")
    try:
        if extension == "csv":
            text = payload.decode("utf-8-sig"); reader = csv.DictReader(StringIO(text))
            headings = [str(h or "").strip().lower() for h in (reader.fieldnames or [])]
            raw_rows = [{str(k or "").strip().lower(): v for k, v in row.items()} for row in reader]
        else:
            wb = load_workbook(BytesIO(payload), read_only=True, data_only=True)
            ws = wb.active; values = ws.iter_rows(values_only=True); first = next(values, ())
            headings = [_normalized(h).lower() for h in first]
            raw_rows = [dict(zip(headings, row)) for row in values]
    except (UnicodeDecodeError, csv.Error, ValueError, OSError) as exc:
        raise HTTPException(422, "The import file is malformed or unreadable.") from exc
    required = {"name", "email", "department"}; missing = sorted(required - set(headings))
    if missing: raise HTTPException(422, f"Missing required columns: {', '.join(missing)}.")
    if len(raw_rows) > 5000: raise HTTPException(422, "An import may contain at most 5,000 data rows.")
    client = get_supabase_client()
    departments = client.table("departments").select("id,name,code").eq("status", "active").execute().data or []
    department_map = {}
    for department in departments:
        for key in (department.get("id"), department.get("name"), department.get("code")):
            if key: department_map[_normalized(key).casefold()] = department["id"]
    existing = {str(row["email"]).casefold() for row in (client.table("supervisors").select("email").execute().data or [])}
    seen = set(); imported = []; errors = []; email_adapter = TypeAdapter(EmailStr)
    for row_number, raw in enumerate(raw_rows, start=2):
        row = {key: _normalized(raw.get(key)) for key in SUPERVISOR_IMPORT_COLUMNS}
        if not any(row.values()): continue
        email = row["email"].lower()
        messages = []
        if any(value.lstrip().startswith(("=", "+", "-", "@")) for value in row.values()):
            messages.append("Spreadsheet formulas are not allowed.")
        if len(row["name"]) < 2: messages.append("Name is required.")
        try: email = str(email_adapter.validate_python(email)).lower()
        except ValidationError: messages.append("Email is invalid.")
        department_id = department_map.get(row["department"].casefold())
        if not department_id: messages.append("Department is not active or was not found.")
        if email in seen: messages.append("Email is duplicated within this file.")
        if email in existing: messages.append("A supervisor with this email already exists.")
        status = row["status"].lower() or "active"
        if status not in {"active", "inactive"}: messages.append("Status must be active or inactive.")
        if messages:
            errors.append({"row": row_number, "email": email, "message": " ".join(messages)}); seen.add(email); continue
        seen.add(email)
        record = {"name": row["name"], "email": email, "phone": row["phone"] or None,
                  "department_id": department_id, "designation": row["designation"] or None,
                  "areas_of_interest": row["areas_of_interest"] or None, "status": status,
                  "approval_status": "approved", "max_students": 30, "created_by": admin["user_id"]}
        try:
            result = client.table("supervisors").insert(record).execute().data or []
            if not result: raise RuntimeError("No record returned")
            imported.append({"row": row_number, "id": result[0]["id"], "email": email}); existing.add(email)
        except Exception:
            errors.append({"row": row_number, "email": email, "message": "The record could not be saved."})
    await write_audit_log("account_created", admin["user_id"], "admin", "supervisor_import",
                          metadata={"filename": file.filename, "imported": len(imported), "failed": len(errors)})
    return {"total_rows": len(imported) + len(errors), "imported_count": len(imported),
            "failed_count": len(errors), "imported": imported, "errors": errors}


@router.get("/exports/{entity}")
async def export(entity: Literal["students", "supervisors"], format: Literal["csv", "xlsx"] = "csv",
                 status: Optional[str] = None, department_id: Optional[str] = None,
                 admin: dict = Depends(require_admin)):
    client = get_supabase_client()
    if entity == "students": fields = "id,name,student_number,registration_number,email,phone,status,created_at,updated_at,departments(name),programs(name),semesters(name),sessions(academic_year),sections(name),projects(title,idea,status,supervisors(name))"
    else: fields = "id,name,email,phone,designation,status,approval_status,email_verified_at,max_students,created_at,updated_at,departments(name),supervisor_program_eligibility(programs(name))"
    q = client.table(entity).select(fields).order("name")
    if status: q = q.eq("status", status)
    if department_id: q = q.eq("department_id", department_id)
    records = q.execute().data or []
    rows = []
    for r in records:
        if entity == "students":
            p = (r.get("projects") or [{}])[0]; rows.append({"Student Name": r.get("name"), "Roll Number": r.get("student_number"), "Registration Number": r.get("registration_number"), "Email": r.get("email"), "Phone": r.get("phone"), "Department": (r.get("departments") or {}).get("name"), "Program": (r.get("programs") or {}).get("name"), "Semester": (r.get("semesters") or {}).get("name"), "Section": (r.get("sections") or {}).get("name"), "Academic Session": (r.get("sessions") or {}).get("academic_year"), "Supervisor": (p.get("supervisors") or {}).get("name"), "Project Title": p.get("title"), "Project Idea": p.get("idea"), "Project Status": p.get("status"), "Account Status": r.get("status"), "Registration Date": r.get("created_at"), "Last Updated": r.get("updated_at")})
        else:
            programs = ", ".join((e.get("programs") or {}).get("name", "") for e in r.get("supervisor_program_eligibility") or []); rows.append({"Supervisor Name": r.get("name"), "Official Email": r.get("email"), "Phone": r.get("phone"), "Department": (r.get("departments") or {}).get("name"), "Programs": programs, "Designation": r.get("designation"), "Approval Status": r.get("approval_status"), "Account Status": r.get("status"), "Email Verified": bool(r.get("email_verified_at")), "Maximum Students": r.get("max_students"), "Registration Date": r.get("created_at"), "Last Updated": r.get("updated_at")})
    columns = list(rows[0]) if rows else (["Student Name"] if entity == "students" else ["Supervisor Name"])
    stamp = datetime.now(timezone.utc).strftime("%Y-%m-%d"); filename = f"{entity}-{stamp}.{format}"
    await write_audit_log("account_updated", admin["user_id"], "admin", "export", metadata={"entity": entity, "format": format, "count": len(rows)})
    return _tabular_response(rows, columns, format, filename.rsplit(".", 1)[0])
